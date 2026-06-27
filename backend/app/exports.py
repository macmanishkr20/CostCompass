"""Boardroom-ready exports: a one-file PDF brief and a multi-sheet workbook.

Both render straight from a validated `Estimation`, so the numbers are exactly
what the engine computed — no recomputation, no drift between the UI and the
downloadable artifacts.
"""

from __future__ import annotations

import io

from .shared.schemas import Estimation


def _money(currency: str, n: float) -> str:
    if isinstance(n, float) and not n.is_integer():
        return f"{currency} {n:,.2f}"
    return f"{currency} {int(n):,}"


# ── PDF (reportlab) ─────────────────────────────────────────────────

def estimation_to_pdf(est: Estimation) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    cur = est.cost_breakdown.currency
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch,
        topMargin=0.75 * inch, bottomMargin=0.75 * inch,
        title=f"{est.project_name} — AI Feasibility & Cost",
    )

    base = getSampleStyleSheet()
    ink = colors.HexColor("#0f172a")
    accent = colors.HexColor("#6366f1")
    muted = colors.HexColor("#64748b")
    h1 = ParagraphStyle("h1", parent=base["Title"], textColor=ink, fontSize=22, spaceAfter=4, alignment=TA_LEFT)
    sub = ParagraphStyle("sub", parent=base["Normal"], textColor=muted, fontSize=10, spaceAfter=10)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], textColor=accent, fontSize=13, spaceBefore=14, spaceAfter=6)
    body = ParagraphStyle("body", parent=base["Normal"], textColor=ink, fontSize=10, leading=15)
    big = ParagraphStyle("big", parent=base["Normal"], textColor=ink, fontSize=15, leading=18)
    small = ParagraphStyle("small", parent=base["Normal"], textColor=muted, fontSize=9, leading=13)

    # Disposition → tone colour for the verdict headline.
    tone_colors = {
        "go": colors.HexColor("#16a34a"),
        "caution": colors.HexColor("#d97706"),
        "stop": colors.HexColor("#dc2626"),
    }

    story: list = []

    story.append(Paragraph(est.project_name, h1))
    mode = "Existing-app enhancement" if est.project_type == "enhancement" else "New build"
    story.append(Paragraph(f"AI Feasibility &amp; Cost Estimate · {mode} · {est.generated_at[:10]}", sub))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0")))

    # The call (canonical verdict) — the single decisive recommendation, willing
    # to say "don't use AI". Optional for estimates persisted before it existed.
    if est.verdict is not None:
        v = est.verdict
        tone = tone_colors.get(v.disposition, ink)
        verdict_style = ParagraphStyle(
            "verdict", parent=base["Normal"], textColor=tone,
            fontSize=18, leading=21, spaceBefore=12, spaceAfter=2,
        )
        story.append(Paragraph(f"<b>{v.headline}</b>", verdict_style))
        story.append(Paragraph(v.one_liner, body))
        if est.confidence is not None:
            cf = est.confidence
            story.append(Spacer(1, 4))
            story.append(Paragraph(
                f"<b>Confidence: {cf.level.capitalize()} ({cf.score}/100)</b> — {cf.rationale}", small,
            ))
            marks = {"positive": "+", "neutral": "·", "negative": "−"}
            for fac in cf.factors:
                story.append(Paragraph(f"{marks.get(fac.impact, '·')} <b>{fac.label}:</b> {fac.detail}", small))

    # Headline.
    f = est.feasibility
    story.append(Paragraph("Recommendation", h2))
    story.append(Paragraph(f"<b>{f.archetype_label}</b> — feasibility {f.score}/100 ({f.rating})", big))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f.rationale, body))

    # Sub-scores.
    s = f.sub_scores
    score_tbl = Table(
        [["AI necessity", "Agentic suitability", "Traditional suitability"],
         [str(s.ai_necessity), str(s.agentic_suitability), str(s.traditional_suitability)]],
        colWidths=[2.2 * inch] * 3,
    )
    score_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
        ("TEXTCOLOR", (0, 0), (-1, 0), muted),
        ("TEXTCOLOR", (0, 1), (-1, 1), ink),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, 1), 18),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    story.append(Spacer(1, 8))
    story.append(score_tbl)

    # Cost summary.
    c = est.cost_breakdown
    story.append(Paragraph("First-year cost", h2))
    cost_rows = [
        ["Line item", "Amount"],
        ["Development (one-off)", _money(cur, c.development.total_cost)],
        ["Infrastructure (annual)", _money(cur, c.infrastructure.annual_cost)],
        ["AI tokens (annual, expected)", _money(cur, c.ai_tokens.annual_cost.expected)],
        ["Maintenance (annual)", _money(cur, c.maintenance.annual_cost)],
        ["Total expected", _money(cur, c.total.expected)],
        ["Total range", f"{_money(cur, c.total.min)} – {_money(cur, c.total.max)}"],
    ]
    cost_tbl = Table(cost_rows, colWidths=[3.6 * inch, 2.8 * inch])
    cost_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -2), 0.5, colors.HexColor("#e2e8f0")),
        ("BACKGROUND", (0, 5), (-1, 5), colors.HexColor("#eef2ff")),
        ("FONTNAME", (0, 5), (-1, 5), "Helvetica-Bold"),
    ]))
    story.append(cost_tbl)

    # ROI.
    r = est.roi_projection
    story.append(Paragraph("Return on investment (3-year)", h2))
    payback = f"{r.payback_months:.1f} months" if r.payback_months is not None else "n/a (run cost exceeds modelled benefit)"
    roi_rows = [
        ["Modelled annual benefit", _money(cur, r.annual_benefit)],
        ["Annual run cost", _money(cur, r.annual_run_cost)],
        ["Net annual benefit", _money(cur, r.net_annual_benefit)],
        ["Payback period", payback],
        ["3-year net value", _money(cur, r.three_year_value)],
        ["3-year ROI", f"{r.roi_percent}%"],
    ]
    roi_tbl = Table(roi_rows, colWidths=[3.6 * inch, 2.8 * inch])
    roi_tbl.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TEXTCOLOR", (0, 0), (0, -1), muted),
        ("TEXTCOLOR", (1, 0), (1, -1), ink),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#eef2ff")),
    ]))
    story.append(roi_tbl)
    if r.assumptions:
        story.append(Paragraph("Benefit basis: " + " ".join(r.assumptions[:2]), small))

    # AI vs Standard.
    cmp = est.comparison
    story.append(Paragraph("AI vs standard approach", h2))
    story.append(Paragraph(cmp.summary, body))
    cmp_rows = [["Approach", "Total (expected)", "Timeline", "Team", "Monthly run"]]
    cmp_rows.append(["AI build", _money(cur, cmp.ai_approach.total_cost.expected), f"{cmp.ai_approach.timeline_weeks} wks", str(cmp.ai_approach.team_size), _money(cur, cmp.ai_approach.monthly_run_cost)])
    cmp_rows.append(["Standard build", _money(cur, cmp.standard_approach.total_cost.expected), f"{cmp.standard_approach.timeline_weeks} wks", str(cmp.standard_approach.team_size), _money(cur, cmp.standard_approach.monthly_run_cost)])
    cmp_tbl = Table(cmp_rows, colWidths=[1.6 * inch, 1.7 * inch, 0.9 * inch, 0.7 * inch, 1.5 * inch])
    cmp_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
        ("TEXTCOLOR", (0, 0), (-1, 0), muted),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("ALIGN", (0, 0), (0, -1), "LEFT"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    story.append(cmp_tbl)

    # Recommendations.
    story.append(Paragraph("Key recommendations", h2))
    for rec in est.recommendations:
        story.append(Paragraph(f"<b>[{rec.priority.upper()}] {rec.title}</b> — {rec.description}", body))
        story.append(Spacer(1, 3))

    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0")))
    story.append(Paragraph(
        "Generated by CostCompass. All figures are computed deterministically; "
        "the language model only classifies use cases.",
        ParagraphStyle("foot", parent=base["Normal"], textColor=muted, fontSize=8, spaceBefore=6),
    ))

    doc.build(story)
    return buf.getvalue()


# ── Excel (openpyxl) ────────────────────────────────────────────────

def estimation_to_excel(est: Estimation) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cur = est.cost_breakdown.currency
    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    accent_font = Font(color="6366F1", bold=True, size=13)
    bold = Font(bold=True)

    def style_header(ws, row: int, ncols: int) -> None:
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

    def autosize(ws, widths: list[int]) -> None:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 1 — Summary.
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = est.project_name
    ws["A1"].font = accent_font
    ws["A2"] = f"AI Feasibility & Cost · {'Enhancement' if est.project_type == 'enhancement' else 'New build'} · {est.generated_at[:10]}"
    f = est.feasibility
    rows: list[tuple[str, object]] = [("Recommendation", f.archetype_label)]
    if est.verdict is not None:
        rows.append(("Verdict", est.verdict.headline))
        rows.append(("Verdict rationale", est.verdict.one_liner))
    if est.confidence is not None:
        rows.append(("Confidence", f"{est.confidence.level.capitalize()} ({est.confidence.score}/100)"))
    rows += [
        ("Feasibility score", f"{f.score}/100 ({f.rating})"),
        ("AI necessity", f.sub_scores.ai_necessity),
        ("Agentic suitability", f.sub_scores.agentic_suitability),
        ("Traditional suitability", f.sub_scores.traditional_suitability),
        ("Total expected (first year)", est.cost_breakdown.total.expected),
        ("Total range", f"{est.cost_breakdown.total.min} – {est.cost_breakdown.total.max}"),
        ("Currency", cur),
    ]
    r = 4
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = bold
        ws.cell(row=r, column=2, value=val)
        r += 1

    # Confidence factors — the deterministic signals behind the score.
    if est.confidence is not None:
        r += 1
        ws.cell(row=r, column=1, value="Confidence factors").font = bold
        ws.cell(row=r, column=2, value=est.confidence.rationale)
        r += 1
        for fac in est.confidence.factors:
            ws.cell(row=r, column=1, value=f"  {fac.label} ({fac.impact})")
            ws.cell(row=r, column=2, value=fac.detail)
            r += 1
    autosize(ws, [30, 60])

    # Sheet 2 — Cost breakdown.
    ws = wb.create_sheet("Cost Breakdown")
    ws.append(["Category", "Detail", f"Amount ({cur})"])
    style_header(ws, 1, 3)
    c = est.cost_breakdown
    for b in c.development.breakdown:
        ws.append(["Development", f"{b.category} ({b.hours}h)", b.cost])
    ws.append(["Development", "Total development", c.development.total_cost])
    for svc in c.infrastructure.services:
        ws.append(["Infrastructure (monthly)", f"{svc.service_name} — {svc.tier}", svc.monthly_cost])
    ws.append(["Infrastructure", "Annual infrastructure", c.infrastructure.annual_cost])
    ws.append(["AI tokens", "Annual (optimistic)", c.ai_tokens.annual_cost.optimistic])
    ws.append(["AI tokens", "Annual (expected)", c.ai_tokens.annual_cost.expected])
    ws.append(["AI tokens", "Annual (pessimistic)", c.ai_tokens.annual_cost.pessimistic])
    ws.append(["Maintenance", f"{c.maintenance.monthly_hours}h/mo @ {_money(cur, c.maintenance.hourly_rate)}/h", c.maintenance.annual_cost])
    ws.append(["TOTAL", "First-year expected", c.total.expected])
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=3).font = bold
    autosize(ws, [26, 44, 18])

    # Sheet 3 — Token model breakdown.
    ws = wb.create_sheet("Tokens")
    ws.append(["Model", "Use cases", "Monthly input tok", "Monthly output tok", f"Monthly cost ({cur})", "In $/1M", "Out $/1M"])
    style_header(ws, 1, 7)
    for m in c.ai_tokens.model_breakdown:
        ws.append([m.model, ", ".join(m.use_cases), m.monthly_input_tokens, m.monthly_output_tokens, m.monthly_cost, m.input_price_per_1m, m.output_price_per_1m])
    autosize(ws, [18, 34, 18, 18, 16, 10, 10])

    # Sheet 4 — ROI.
    ws = wb.create_sheet("ROI")
    roi = est.roi_projection
    ws.append(["Metric", "Value"])
    style_header(ws, 1, 2)
    ws.append(["Annual benefit", roi.annual_benefit])
    ws.append(["Annual run cost", roi.annual_run_cost])
    ws.append(["Net annual benefit", roi.net_annual_benefit])
    ws.append(["Payback (months)", roi.payback_months if roi.payback_months is not None else "n/a"])
    ws.append(["3-year net value", roi.three_year_value])
    ws.append(["3-year ROI %", roi.roi_percent])
    ws.append([])
    ws.append(["Use case", "Minutes/call", "Loaded $/hr", "Automation %", "Value/call", "Annual calls", "Annual value"])
    style_header(ws, ws.max_row, 7)
    for d in roi.value_drivers:
        ws.append([
            d.use_case,
            d.minutes_per_call if d.minutes_per_call is not None else "override",
            d.loaded_hourly_rate if d.loaded_hourly_rate is not None else "—",
            d.automation_rate_percent if d.automation_rate_percent is not None else "—",
            d.value_per_call,
            d.annual_calls,
            d.annual_value,
        ])
    ws.append([])
    ws.append(["Basis"])
    style_header(ws, ws.max_row, 1)
    for a in roi.assumptions:
        ws.append([a])
    ws.append([])
    ws.append(["Month", "Cumulative net"])
    style_header(ws, ws.max_row, 2)
    for pt in roi.curve:
        ws.append([pt.month, pt.cumulative_net])
    autosize(ws, [28, 13, 12, 14, 12, 14, 16])

    # Sheet 5 — Recommendations.
    ws = wb.create_sheet("Recommendations")
    ws.append(["Priority", "Category", "Title", "Description", "Estimated impact"])
    style_header(ws, 1, 5)
    for rec in est.recommendations:
        ws.append([rec.priority, rec.category, rec.title, rec.description, rec.estimated_impact])
    for col in range(1, 6):
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws, [10, 14, 32, 50, 40])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
